"""Export separate Vietnamese annotation workbooks with opaque, shuffled IDs.

Requires openpyxl (an optional export dependency). The coordinator mapping is kept
outside the reviewer ZIP. Never marks records reviewed or changes the benchmark.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import random
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
BLUE = '17365D'
YELLOW = 'FFF2CC'
GREY = 'F2F4F7'
GREEN = 'E2F0D9'
ORANGE = 'FCE4D6'

QUESTION_COLUMNS = [
    ('benchmark_id','Mã câu',20,False),
    ('query','Câu hỏi',65,False),
    ('behavior_label','Nhãn hành vi',22,True),
    ('question_decision','Giữ / sửa / loại câu',23,True),
    ('anchor_evidence','Đoạn bằng chứng gợi ý',85,False),
    ('source_url','Link nguồn',38,False),
    ('reference_answer','Câu trả lời tham chiếu',55,True),
    ('query_revision','Câu hỏi viết lại (nếu cần)',55,True),
    ('task_type','Loại nhu cầu du lịch',28,True),
    ('evidence_search_notes','Ghi chú tìm / kiểm tra bằng chứng',60,True),
    ('reviewer_id','Mã người thẩm định',22,True),
    ('review_status','Trạng thái',22,True),
    ('reviewer_notes','Ghi chú khác / lý do loại',50,True),
]
QREL_COLUMNS = [
    ('benchmark_id','Mã câu',20,False),
    ('query','Câu hỏi',65,False),
    ('chunk_id','Mã đoạn văn',25,False),
    ('relevance_grade','Mức bằng chứng',22,True),
    ('candidate_evidence','Đoạn văn cần đánh giá',85,False),
    ('evidence_quote','Trích nguyên văn làm bằng chứng',60,True),
    ('source_url','Link nguồn',38,False),
    ('reviewer_id','Mã người thẩm định',22,True),
    ('review_status','Trạng thái',22,True),
    ('reviewer_notes','Ghi chú đánh giá',50,True),
]


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as stream:
        return list(csv.DictReader(stream))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def literal(cell,value):
    """Corpus/reviewer text must remain text, even when starting with '='."""
    cell.value=str(value or '')
    cell.data_type='s'


def dropdown(sheet,column,last_row,values,message):
    rule=DataValidation(type='list',formula1='"'+','.join(values)+'"',allow_blank=True)
    rule.errorTitle='Giá trị chưa hợp lệ'
    rule.error='Chọn một giá trị trong danh sách, hoặc để trống khi chưa đánh giá.'
    rule.showErrorMessage=True
    rule.errorStyle='stop'
    rule.promptTitle='Hướng dẫn điền'
    rule.prompt=message[:255]
    rule.showInputMessage=True
    rule.showDropDown=False
    sheet.add_data_validation(rule)
    rule.add(f'{column}2:{column}{last_row}')


def data_sheet(workbook,title,columns,rows,table_name):
    sheet=workbook.create_sheet(title)
    for col,(key,label,width,editable) in enumerate(columns,1):
        cell=sheet.cell(1,col,label)
        cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor=BLUE)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        cell.comment=Comment(f'Trường: {key}\n'+('Ô vàng: người thẩm định điền.' if editable else 'Thông tin gốc: giữ nguyên; nếu cần sửa câu, dùng cột câu hỏi viết lại.'),'Hướng dẫn')
        sheet.column_dimensions[get_column_letter(col)].width=width
    for index,row in enumerate(rows,2):
        for col,(key,_,_,editable) in enumerate(columns,1):
            cell=sheet.cell(index,col)
            literal(cell,row.get(key,''))
            cell.font=Font(name='Calibri',size=11)
            cell.fill=PatternFill('solid',fgColor=YELLOW if editable else GREY)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if key=='source_url' and cell.value.startswith(('https://','http://')):
                cell.hyperlink=cell.value
                cell.font=Font(name='Calibri',size=11,color='0563C1',underline='single')
        lines=max(sum(max(1,math.ceil(len(part)/max(1,width-3)))
                      for part in str(row.get(key,'')).split('\n'))
                  for key,_,width,_ in columns)
        sheet.row_dimensions[index].height=min(360,max(95,lines*16+12))
    sheet.row_dimensions[1].height=36
    sheet.freeze_panes='B2'
    sheet.sheet_view.zoomScale=80
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    sheet.print_title_rows='1:1'
    ref=f'A1:{get_column_letter(len(columns))}{len(rows)+1}'
    table=Table(displayName=table_name,ref=ref)
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
    sheet.add_table(table)
    return sheet


def guide(workbook,reviewer,questions,qrels):
    sheet=workbook.active
    sheet.title='HUONG_DAN'
    sheet.column_dimensions['A'].width=33
    sheet.column_dimensions['B'].width=115
    entries=[
        ('BỘ THẨM ĐỊNH V2',f'File riêng cho người thẩm định số {reviewer}. Có {questions} câu hỏi nháp và {qrels} cặp câu hỏi–đoạn văn. Đây là dữ liệu chờ đánh giá, chưa phải bộ nhãn đã xác nhận.'),
        ('Cách bắt đầu','1) Đọc hướng dẫn. 2) Làm thử 5–10 câu. 3) Điền các ô màu vàng trong CAU_HOI và BANG_CHUNG. 4) Lưu lại file .xlsx và gửi bản đã điền cho người điều phối.'),
        ('Đánh giá độc lập','Mỗi người chỉ nhận file của mình. Không xem kết quả người khác trước khi hoàn tất lượt đánh giá độc lập. Mỗi người tự điền một mã riêng, ví dụ R01; để trống thông tin chưa biết, không dùng tên người khác.'),
        ('Giữ nguyên thông tin gốc','Không đổi Mã câu, Mã đoạn văn, câu hỏi gốc, đoạn nguồn hoặc xóa/sắp xếp riêng từng cột. Có thể lọc bảng bằng nút ở tiêu đề. Nếu cần sửa câu, ghi vào “Câu hỏi viết lại”.'),
        ('Ô màu vàng','Đây là các trường bạn điền. Cột nhãn, quyết định, loại nhu cầu và trạng thái có danh sách chọn. Câu trả lời tham chiếu là đáp án đúng theo bằng chứng, không phải câu trả lời của mô hình.'),
        ('answer','Câu hỏi có thể được trả lời bằng bằng chứng trong corpus và không cần sửa tiền đề sai. Ghi đáp án ngắn, đúng phạm vi nguồn và thời gian.'),
        ('correct','Câu hỏi chứa tiền đề sai có thể bác bỏ bằng bằng chứng. Ghi rõ chỗ sai, thông tin đúng và trích dẫn hỗ trợ; không chỉ trả lời “sai”.'),
        ('refuse','Chỉ dùng khi đã kiểm tra trong phạm vi corpus mà vẫn không có đủ bằng chứng. Một đoạn gợi ý không nêu thông tin KHÔNG đủ để kết luận refuse. Không suy đoán rằng thông tin không tồn tại ngoài đời.'),
        ('unclear','Chưa đủ cơ sở phân loại, thông tin mâu thuẫn, câu quá mơ hồ hoặc chưa kiểm tra đủ bằng chứng. Ghi lý do và chọn trạng thái NEEDS_REVIEW.'),
        ('KEEP / REVISE / REJECT / UNSURE','KEEP: giữ câu. REVISE: cần viết lại. REJECT: không thích hợp, ví dụ hỏi footer hoặc không có giá trị cho nhiệm vụ du lịch. UNSURE: cần thảo luận thêm. Không xóa dòng bị loại; ghi lý do.'),
        ('Mức bằng chứng 0','Đoạn không liên quan hoặc không hữu ích cho thông tin đang hỏi.'),
        ('Mức bằng chứng 1','Đoạn có liên quan nhưng chưa đủ để trả lời thông tin cụ thể hoặc sửa tiền đề. Không coi liên quan chủ đề là đủ bằng chứng.'),
        ('Mức bằng chứng 2','Đoạn đủ hỗ trợ câu trả lời hoặc sửa tiền đề. Với mức 1 hoặc 2, chép một trích đoạn nguyên văn phù hợp vào cột bằng chứng; không tự viết lại thành câu nguồn.'),
        ('Một câu có nhiều đoạn','Mỗi dòng BANG_CHUNG là một cặp câu hỏi–đoạn văn. Có thể nhiều đoạn cùng được mức 2. Có thể lọc theo Mã câu để xem các đoạn ứng viên của cùng câu hỏi.'),
        ('TODO / DONE / NEEDS_REVIEW','TODO: chưa làm. DONE: bạn đã hoàn tất đánh giá riêng của mình. NEEDS_REVIEW: còn vướng hoặc cần kiểm tra thêm. DONE không tự động trở thành nhãn VERIFIED của nghiên cứu.'),
        ('Loại nhu cầu du lịch','destination_choice: chọn điểm đến; visit_timing: thời điểm tham quan; price_eligibility: giá/điều kiện ưu đãi; access_logistics: đi lại/tiếp cận; cultural_event_information: thông tin văn hóa/sự kiện; other: khác.'),
        ('Giới hạn của bộ hiện tại','Đây là bản nháp tự động. Có câu dài, thiếu tự nhiên, thiên về số liệu hoặc lẫn metadata. Hãy đánh dấu sửa/loại khi cần. Đoạn gợi ý không mặc nhiên là đáp án đúng.'),
        ('Chưa đủ pool bằng chứng','File chứa đoạn nguồn và các ứng viên truy xuất đã chuẩn bị; chưa có đầy đủ dense/RRF cho mọi câu. Không buộc phải chốt refuse. Hãy ghi những nguồn/từ khóa đã kiểm tra và dùng NEEDS_REVIEW khi chưa đủ.'),
        ('Thông tin theo thời gian','Phân biệt “theo bài nguồn năm X” với tình hình hiện tại. Không lấy tin cũ để khẳng định giá vé, lịch mở cửa hoặc chính sách đang có hiệu lực hôm nay.'),
        ('Ví dụ minh họa','Nếu nguồn ghi “bờ biển dài 116 km”: câu hỏi chiều dài → answer; câu khẳng định “160 km” → correct. Câu hỏi ưu đãi sinh viên chưa được nguồn đề cập → chưa chốt refuse nếu chưa tìm thêm.'),
        ('Mã câu trung tính','Mã câu đã đổi và thứ tự đã xáo trộn để tránh lộ nhãn tự động đề xuất. Hai người vẫn đánh giá cùng một tập câu. Bảng đối chiếu mã do người điều phối giữ riêng.'),
        ('Cách gửi kết quả','Giữ định dạng Excel .xlsx và gửi lại file riêng đã điền. Hai file được ghép để tìm bất đồng; người thứ ba phân xử sau lượt đánh giá độc lập. Không trao đổi nhãn trước đó.'),
    ]
    for row,(heading,content) in enumerate(entries,1):
        literal(sheet.cell(row,1),heading)
        literal(sheet.cell(row,2),content)
        for cell in sheet[row]:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            cell.font=Font(name='Calibri',size=12,bold=cell.column==1)
        sheet.row_dimensions[row].height=65 if row!=1 else 75
    for cell in sheet[1]:
        cell.fill=PatternFill('solid',fgColor=BLUE)
        cell.font=Font(name='Calibri',size=15,bold=True,color='FFFFFF')
    sheet.freeze_panes='B2'
    sheet.sheet_view.zoomScale=90


def progress(workbook,question_count,qrel_count):
    sheet=workbook.create_sheet('TIEN_DO')
    sheet.append(['Nội dung','Số lượng'])
    sheet.append(['Tổng câu hỏi',question_count])
    sheet.append(['Câu hỏi đã chọn DONE',f'=COUNTIF(CAU_HOI!L2:L{question_count+1},"DONE")'])
    sheet.append(['Câu hỏi cần xem lại',f'=COUNTIF(CAU_HOI!L2:L{question_count+1},"NEEDS_REVIEW")'])
    sheet.append(['Tổng cặp câu–đoạn',qrel_count])
    sheet.append(['Cặp câu–đoạn đã chọn DONE',f'=COUNTIF(BANG_CHUNG!I2:I{qrel_count+1},"DONE")'])
    sheet.append(['Cặp câu–đoạn cần xem lại',f'=COUNTIF(BANG_CHUNG!I2:I{qrel_count+1},"NEEDS_REVIEW")'])
    sheet.append(['Lưu ý','Tiến độ là trạng thái người đánh giá tự chọn; không xác nhận nhãn đã đúng hoặc sẵn sàng công bố.'])
    sheet.column_dimensions['A'].width=65
    sheet.column_dimensions['B'].width=65
    for row in sheet:
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[row[0].row].height=32
    workbook.calculation.fullCalcOnLoad=True


def make_workbook(path,reviewer,questions,qrels,mapping):
    wb=Workbook()
    guide(wb,reviewer,len(questions),len(qrels))
    qrows=[dict(row,benchmark_id=mapping[row['benchmark_id']]) for row in questions]
    erows=[dict(row,benchmark_id=mapping[row['benchmark_id']]) for row in qrels]
    random.Random(20261031+reviewer).shuffle(qrows)
    random.Random(20261131+reviewer).shuffle(erows)
    qs=data_sheet(wb,'CAU_HOI',QUESTION_COLUMNS,qrows,'CauHoi')
    es=data_sheet(wb,'BANG_CHUNG',QREL_COLUMNS,erows,'BangChung')
    dropdown(qs,'C',len(qrows)+1,['answer','correct','refuse','unclear'],'Chọn theo bằng chứng; nếu chưa đủ cơ sở, chọn unclear.')
    dropdown(qs,'D',len(qrows)+1,['KEEP','REVISE','REJECT','UNSURE'],'Đánh giá chất lượng câu; không xóa dòng bị loại.')
    dropdown(qs,'I',len(qrows)+1,['destination_choice','visit_timing','price_eligibility','access_logistics','cultural_event_information','other'],'Chọn loại nhu cầu phù hợp nhất.')
    dropdown(qs,'L',len(qrows)+1,['TODO','DONE','NEEDS_REVIEW'],'Chọn DONE khi hoàn tất lượt đánh giá của riêng bạn.')
    dropdown(es,'D',len(erows)+1,['0','1','2','unclear'],'0: không hữu ích; 1: liên quan nhưng chưa đủ; 2: đủ bằng chứng; unclear: chưa rõ.')
    dropdown(es,'I',len(erows)+1,['TODO','DONE','NEEDS_REVIEW'],'Mỗi dòng là một câu hỏi và một đoạn ứng viên.')
    for sheet,column,total in [(qs,'L',len(qrows)),(es,'I',len(erows))]:
        for label,color in [('DONE',GREEN),('NEEDS_REVIEW',ORANGE)]:
            sheet.conditional_formatting.add(f'{column}2:{column}{total+1}',
                CellIsRule(operator='equal',formula=['"'+label+'"'],fill=PatternFill('solid',fgColor=color)))
    progress(wb,len(qrows),len(erows))
    wb.properties.title='Thẩm định benchmark du lịch V2'
    wb.properties.subject='Phiếu độc lập; nhãn chờ người đánh giá'
    wb.properties.creator='Nhóm nghiên cứu — mẫu chuẩn bị bằng công cụ'
    wb.save(path)


def check_workbook(path,expected_questions,expected_qrels,original_ids):
    wb=load_workbook(path,data_only=False)
    assert wb['CAU_HOI'].max_row==expected_questions+1
    assert wb['BANG_CHUNG'].max_row==expected_qrels+1
    qids={wb['CAU_HOI'].cell(i,1).value for i in range(2,expected_questions+2)}
    assert len(qids)==expected_questions
    assert all(wb['BANG_CHUNG'].cell(i,1).value in qids for i in range(2,expected_qrels+2))
    for name,columns,count in [('CAU_HOI',[3,4,7,8,9,10,11,13],expected_questions),
                               ('BANG_CHUNG',[4,6,8,10],expected_qrels)]:
        for index in range(2,count+2):
            assert all(wb[name].cell(index,col).value is None for col in columns)
    assert len(wb['CAU_HOI'].data_validations.dataValidation)==4
    assert len(wb['BANG_CHUNG'].data_validations.dataValidation)==2
    assert all(sheet.sheet_state=='visible' for sheet in wb)
    assert all(cell.data_type!='f' for sheet in wb if sheet.title!='TIEN_DO' for row in sheet for cell in row)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                assert cell.value not in original_ids if isinstance(cell.value,str) else True
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith('.xml'):
                content=archive.read(name)
                assert not any(qid.encode() in content for qid in original_ids)
    return qids


def export(packet_dir,out):
    if out.exists():
        raise FileExistsError('Choose a new export directory; existing files are retained: '+str(out))
    questions=[read_csv(packet_dir/f'reviewer_{n}_questions.csv') for n in (1,2)]
    qrels=[read_csv(packet_dir/f'reviewer_{n}_qrels.csv') for n in (1,2)]
    ids={q['benchmark_id'] for q in questions[0]}
    if not ids or ids!={q['benchmark_id'] for q in questions[1]}:
        raise ValueError('Reviewer question sets differ or are empty')
    if any(q['behavior_label'] or q['reviewer_id'] or q['review_status']!='TODO' for rows in questions for q in rows):
        raise ValueError('Expected unreviewed source packets; refusing to drop existing reviewer judgments')
    if any(r['relevance_grade'] or r['reviewer_id'] or r['review_status']!='TODO' for rows in qrels for r in rows):
        raise ValueError('Expected unreviewed qrels; refusing to drop existing reviewer judgments')
    mapping={qid:'Q-'+hashlib.sha256(('review-v2:'+qid).encode()).hexdigest()[:12].upper() for qid in ids}
    if len(set(mapping.values()))!=len(mapping):
        raise ValueError('Opaque ID collision')
    out.mkdir(parents=True)
    files=[]
    check_ids=[]
    for reviewer in (1,2):
        path=out/f'Nguoi_tham_dinh_{reviewer:02d}.xlsx'
        make_workbook(path,reviewer,questions[reviewer-1],qrels[reviewer-1],mapping)
        check_ids.append(check_workbook(path,len(questions[reviewer-1]),len(qrels[reviewer-1]),ids))
        files.append(path)
    assert check_ids[0]==check_ids[1]
    mapping_path=out/'DIEU_PHOI_doi_chieu_ma_GIU_RIENG.csv'
    with mapping_path.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=['review_id','benchmark_id'])
        writer.writeheader()
        writer.writerows(dict(review_id=mapping[qid],benchmark_id=qid) for qid in sorted(ids))
    instructions=out/'DOC_TRUOC_KHI_GUI.txt'
    instructions.write_text(
        'BỘ THẨM ĐỊNH V2\n\n'
        'Gửi Nguoi_tham_dinh_01.xlsx cho người thứ nhất; gửi Nguoi_tham_dinh_02.xlsx cho người thứ hai.\n'
        'Mỗi người chỉ xem file của mình và đọc tab HUONG_DAN trước khi điền ô vàng.\n'
        f'Đây là {len(questions[0])} câu nháp và {len(qrels[0])} cặp bằng chứng ban đầu, chưa phải nhãn đã xác nhận.\n'
        'Không trao đổi nhãn trước khi hoàn tất đánh giá độc lập.\n'
        'Chưa đủ bằng chứng: chọn unclear / NEEDS_REVIEW và ghi lý do.\n'
        'Lưu và gửi lại file .xlsx đã điền. Không xóa dòng hoặc đổi mã câu/đoạn.\n'
        'Bảng đối chiếu mã do điều phối giữ riêng, không nằm trong gói ZIP cho người thẩm định.\n',encoding='utf-8-sig')
    zip_path=out/'Bo_Excel_tham_dinh_V2.zip'
    with zipfile.ZipFile(zip_path,'x',compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files+[instructions]:
            archive.write(path,path.name)
    with zipfile.ZipFile(zip_path) as archive:
        assert archive.testzip() is None
        assert set(archive.namelist())=={p.name for p in files+[instructions]}
    manifest=dict(questions_per_reviewer=len(questions[0]),passage_reviews_per_reviewer=len(qrels[0]),
                  source_packet_manifest_sha256=sha(packet_dir/'packet_manifest.json'),
                  source_hashes={p.name:sha(p) for p in packet_dir.glob('*.csv')},
                  files={p.name:dict(bytes=p.stat().st_size,sha256=sha(p)) for p in files+[zip_path]},
                  coordinator_mapping=mapping_path.name,labels_prefilled=False,
                  reviewer_ids_prefilled=False,opaque_ids=True,shuffled_order=True,
                  coordinator_mapping_in_reviewer_zip=False)
    (out/'export_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--packets',type=Path,default=ROOT/'data/research_v2/reviewer_packets_002')
    parser.add_argument('--out',type=Path,default=ROOT/'data/research_v2/excel_exports_001')
    args=parser.parse_args()
    export(args.packets,args.out)
